"""Create a Czech PowerPoint presentation for the spatial index semester project."""

from __future__ import annotations

import csv
from pathlib import Path

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
PRESENTATION = ROOT / "presentation"
ASSETS = PRESENTATION / "assets"
OUTPUT = PRESENTATION / "MH_semestralka_p__index_prezentace.pptx"

BG = RGBColor(248, 249, 251)
DARK = RGBColor(32, 40, 48)
MUTED = RGBColor(95, 105, 115)
BLUE = RGBColor(44, 100, 170)
GREEN = RGBColor(35, 135, 92)
ORANGE = RGBColor(210, 125, 45)
RED = RGBColor(180, 70, 70)
LIGHT = RGBColor(226, 232, 240)

METHOD_LABELS = {
    "linear_search": "bez indexu",
    "tile_index": "dlaždicový index",
    "rtree_index": "R-tree / STRtree",
    "linear_search_radius": "bez indexu",
    "tile_index_radius": "dlaždicový index",
    "rtree_index_radius": "R-tree / STRtree",
    "linear_search_polygon": "bez indexu",
    "tile_index_polygon": "dlaždicový index",
    "rtree_index_polygon": "R-tree / STRtree",
}

METHOD_COLORS = {
    "bez indexu": "#b84a4a",
    "dlaždicový index": "#23875c",
    "R-tree / STRtree": "#2c64aa",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    sample = path.read_text(encoding="utf-8-sig")[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def fnum(row: dict[str, str], key: str) -> float:
    value = row[key]
    return float(value) if value else 0.0


def make_grouped_bar(
    rows: list[dict[str, str]],
    x_key: str,
    title: str,
    output: Path,
    y_key: str = "total_seconds",
    log_y: bool = False,
) -> None:
    methods = []
    x_values = []
    data: dict[tuple[str, str], float] = {}

    for row in rows:
        method = METHOD_LABELS[row["method"]]
        x_value = row[x_key]
        if method not in methods:
            methods.append(method)
        if x_value not in x_values:
            x_values.append(x_value)
        data[(x_value, method)] = fnum(row, y_key)

    x = list(range(len(x_values)))
    width = 0.24
    fig, ax = plt.subplots(figsize=(8.2, 4.6), dpi=180)

    for index, method in enumerate(methods):
        shift = (index - (len(methods) - 1) / 2) * width
        values = [data.get((value, method), 0.0) for value in x_values]
        ax.bar(
            [item + shift for item in x],
            values,
            width=width,
            label=method,
            color=METHOD_COLORS[method],
        )

    ax.set_title(title, fontsize=13, pad=10)
    ax.set_ylabel("celkový čas [s]")
    ax.set_xticks(x)
    ax.set_xticklabels(x_values)
    if log_y:
        ax.set_yscale("log")
        ax.set_ylabel("celkový čas [s], log měřítko")
    ax.grid(axis="y", alpha=0.22)
    ax.legend(frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(output, bbox_inches="tight")
    plt.close(fig)


def make_polygon_counts_chart(rows: list[dict[str, str]], output: Path) -> None:
    filtered = [
        row
        for row in rows
        if row["dataset_size"] == "50000" and row["method"] == "linear_search_polygon"
    ]
    filtered.sort(key=lambda item: int(item["found_points"]), reverse=True)

    names = [row["kraj_name"] for row in filtered]
    counts = [int(row["found_points"]) for row in filtered]

    fig, ax = plt.subplots(figsize=(8.4, 5.2), dpi=180)
    ax.barh(names, counts, color="#2c64aa")
    ax.invert_yaxis()
    ax.set_title("Počet bodů v krajích, dataset 50 000 bodů", fontsize=13, pad=10)
    ax.set_xlabel("počet bodů")
    ax.grid(axis="x", alpha=0.22)
    ax.tick_params(axis="y", labelsize=8)
    fig.tight_layout()
    fig.savefig(output, bbox_inches="tight")
    plt.close(fig)


def make_charts() -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)

    search_ok = read_csv(RESULTS / "benchmark_search.csv")
    search_cr = read_csv(RESULTS / "benchmark_search_cr.csv")
    radius_cr = read_csv(RESULTS / "benchmark_radius_query_cr.csv")
    polygon_cr = read_csv(RESULTS / "benchmark_polygon_query_cr.csv")
    polygon_counts = read_csv(RESULTS / "polygon_query_counts_by_kraj.csv")

    charts = {
        "olomouc": ASSETS / "benchmark_olomoucky_kraj.png",
        "cr": ASSETS / "benchmark_ceska_republika.png",
        "radius": ASSETS / "benchmark_radius_query.png",
        "polygon": ASSETS / "benchmark_polygon_query.png",
        "counts": ASSETS / "polygon_counts_by_kraj.png",
    }

    make_grouped_bar(
        [row for row in search_ok if row["dataset_size"] == "5000"],
        "query_window_size_m",
        "Olomoucký kraj: 5 000 bodů, 1 000 dotazů",
        charts["olomouc"],
        log_y=True,
    )
    make_grouped_bar(
        [row for row in search_cr if row["dataset_size"] == "50000"],
        "query_window_size_m",
        "Česká republika: 50 000 bodů, 1 000 dotazů",
        charts["cr"],
        log_y=True,
    )
    make_grouped_bar(
        [row for row in radius_cr if row["dataset_size"] == "50000"],
        "radius_m",
        "Radius query: ČR, 50 000 bodů",
        charts["radius"],
        log_y=True,
    )
    make_grouped_bar(
        polygon_cr,
        "dataset_size",
        "Polygon query: 14 krajů ČR",
        charts["polygon"],
        log_y=False,
    )
    make_polygon_counts_chart(polygon_counts, charts["counts"])

    return charts


def set_slide_bg(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    left = Inches(0.55)
    top = Inches(0.35)
    width = Inches(12.2)
    height = Inches(0.72)
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    para = frame.paragraphs[0]
    para.text = title
    para.font.size = Pt(25)
    para.font.bold = True
    para.font.color.rgb = DARK

    if subtitle:
        sub = slide.shapes.add_textbox(left, Inches(1.02), width, Inches(0.38))
        sub_frame = sub.text_frame
        sub_frame.clear()
        p = sub_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = MUTED


def add_footer(slide, number: int) -> None:
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(7.15), Inches(12.2), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT
    line.line.color.rgb = LIGHT

    box = slide.shapes.add_textbox(Inches(10.6), Inches(7.18), Inches(2.1), Inches(0.25))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = f"MH_semestralka_p__index | {number}"
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(8)
    p.font.color.rgb = MUTED


def add_bullets(slide, bullets: list[str], left=0.75, top=1.55, width=5.9, size=18) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.2))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.clear()

    for index, text in enumerate(bullets):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = text
        para.level = 0
        para.font.size = Pt(size)
        para.font.color.rgb = DARK
        para.space_after = Pt(10)
        para.line_spacing = 1.08


def add_note(slide, text: str, left=7.1, top=1.6, width=5.3, height=4.9) -> None:
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = LIGHT
    shape.line.width = Pt(1)

    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.24)
    frame.margin_right = Inches(0.24)
    frame.margin_top = Inches(0.18)
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BLUE
    p.alignment = PP_ALIGN.CENTER


def add_chart(slide, image: Path, left=6.75, top=1.55, width=5.9) -> None:
    slide.shapes.add_picture(str(image), Inches(left), Inches(top), width=Inches(width))


def add_process(slide, labels: list[str], top=2.25) -> None:
    colors = [RED, ORANGE, GREEN, BLUE]
    width = 2.45
    gap = 0.55
    left = 0.75
    for index, label in enumerate(labels):
        x = left + index * (width + gap)
        rect = slide.shapes.add_shape(1, Inches(x), Inches(top), Inches(width), Inches(1.1))
        rect.fill.solid()
        rect.fill.fore_color.rgb = colors[index % len(colors)]
        rect.line.color.rgb = colors[index % len(colors)]
        frame = rect.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        if index < len(labels) - 1:
            arrow = slide.shapes.add_textbox(Inches(x + width + 0.08), Inches(top + 0.33), Inches(0.4), Inches(0.3))
            p2 = arrow.text_frame.paragraphs[0]
            p2.text = ">"
            p2.font.size = Pt(22)
            p2.font.bold = True
            p2.font.color.rgb = MUTED


def add_metric_row(slide, metrics: list[tuple[str, str]], top=5.45) -> None:
    left = 0.75
    width = 3.0
    gap = 0.28
    for index, (label, value) in enumerate(metrics):
        box = slide.shapes.add_shape(1, Inches(left + index * (width + gap)), Inches(top), Inches(width), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = LIGHT
        frame = box.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        p.text = value
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = BLUE
        p.alignment = PP_ALIGN.CENTER
        p2 = frame.add_paragraph()
        p2.text = label
        p2.font.size = Pt(9)
        p2.font.color.rgb = MUTED
        p2.alignment = PP_ALIGN.CENTER


def build_presentation(charts: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title: str, subtitle: str | None = None):
        s = prs.slides.add_slide(blank)
        set_slide_bg(s)
        add_title(s, title, subtitle)
        add_footer(s, len(prs.slides))
        return s

    s = slide("Prostorový index pro vyhledávání bodů", "Semestrální práce | MH_semestralka_p__index")
    add_bullets(
        s,
        [
            "Porovnání vyhledávání bez indexu, vlastního dlaždicového indexu a knihovního R-tree / STRtree.",
            "Testy nad body v Olomouckém kraji i nad celou Českou republikou.",
            "Důraz na praktický výkon u bbox, radius a polygon query.",
        ],
        width=8.2,
        size=20,
    )
    add_note(s, "Hlavní otázka:\nKdy prostorový index skutečně zrychlí dotazy?", left=9.0, top=2.1, width=3.4, height=2.5)

    s = slide("Cíl práce")
    add_bullets(
        s,
        [
            "Vytvořit a otestovat jednoduchý prostorový index nad bodovými daty.",
            "Porovnat ho s lineárním průchodem a s knihovním indexem GeoPandas/Shapely.",
            "Ověřit správnost výsledků: všechny metody musí vracet stejné počty bodů.",
            "Vyhodnotit chování při různých velikostech datasetu a dotazu.",
        ],
        width=7.2,
    )
    add_note(s, "Výstupem není jen kód, ale měřitelný rozdíl v časech dotazů.", left=8.45, top=2.0, width=3.8, height=2.5)

    s = slide("Použitá data")
    add_bullets(
        s,
        [
            "Polygon Olomouckého kraje a polygon celé České republiky.",
            "Vrstva 14 krajů ČR pro polygon query.",
            "Náhodné body generované uvnitř skutečných polygonů.",
            "Datasety: Olomoucký kraj 500, 1 000, 5 000 bodů; ČR 5 000, 10 000, 50 000 bodů.",
            "Zdroj hranic: veřejná ArcGIS REST služba České geologické služby.",
        ],
        width=7.4,
        size=17,
    )
    add_metric_row(s, [("krajů v polygon query", "14"), ("největší dataset", "50 000"), ("dotazů v benchmarku", "1 000")])

    s = slide("Proč EPSG:5514")
    add_bullets(
        s,
        [
            "EPSG:4326 je vhodný výměnný formát pro GIS a webové mapy.",
            "Pro výpočty je potřeba pracovat v metrech, ne ve stupních.",
            "EPSG:5514 (S-JTSK / Křovák East North) umožňuje přímo používat velikost dlaždice, okna i radius v metrech.",
            "Díky tomu je benchmark interpretovatelný: 5 000 znamená 5 km.",
        ],
        width=7.4,
        size=17,
    )
    add_note(s, "Pracovní CRS:\nEPSG:5514\n\nDotazy:\nmetry, bboxy, vzdálenosti", left=8.2, top=1.75, width=3.9, height=3.25)

    s = slide("Vyhledávání bez indexu")
    add_process(s, ["dotaz", "projdi všechny body", "test bbox / vzdálenost", "výsledek"])
    add_bullets(
        s,
        [
            "Lineární přístup je jednoduchý a má nulový čas stavby indexu.",
            "Každý dotaz ale kontroluje celý dataset.",
            "S rostoucím počtem bodů roste čas téměř přímo úměrně.",
        ],
        top=4.25,
        width=10.8,
        size=18,
    )

    s = slide("Dlaždicový prostorový index")
    add_process(s, ["rozděl prostor", "ulož body do dlaždic", "vyber dotčené dlaždice", "ověř kandidáty"])
    add_bullets(
        s,
        [
            "Vlastní implementace ukládá body podle dvojice celočíselných souřadnic dlaždice.",
            "Dotaz se převede na rozsah dlaždic, které protínají bbox.",
            "Přesný test se dělá jen nad kandidáty z těchto dlaždic.",
            "V projektu: 1 000 m pro Olomoucký kraj, 5 000 m pro ČR.",
        ],
        top=4.05,
        width=11.2,
        size=17,
    )

    s = slide("R-tree / STRtree index")
    add_process(s, ["geometrie", "sindex", "bbox kandidáti", "přesné ověření"])
    add_bullets(
        s,
        [
            "GeoPandas/Shapely poskytuje knihovní prostorový index přes `gdf.sindex`.",
            "Index organizuje bounding boxy geometrií tak, aby se rychle našli kandidáti.",
            "Stejně jako u dlaždic není bbox kandidát automaticky finální výsledek.",
            "Výhoda: obecnější řešení pro různé geometrie; nevýhoda: režie Python/GeoPandas v malých dotazech.",
        ],
        top=4.05,
        width=11.2,
        size=17,
    )

    s = slide("Benchmark: Olomoucký kraj", "Bbox query, dataset 5 000 bodů, 1 000 dotazů")
    add_bullets(
        s,
        [
            "Dlaždicový index je nejrychlejší ve všech velikostech dotazovacího okna.",
            "Lineární průchod je použitelný u malých dat, ale s počtem bodů rychle roste.",
            "R-tree/STRtree má v tomto benchmarku vysokou režii a pro malé bodové dotazy nevychází výhodně.",
        ],
        width=5.5,
        size=16,
    )
    add_chart(s, charts["olomouc"])

    s = slide("Benchmark: Česká republika", "Bbox query, dataset 50 000 bodů, 1 000 dotazů")
    add_bullets(
        s,
        [
            "Na větším území a větším datasetu je rozdíl mezi lineárním průchodem a indexem výraznější.",
            "Dlaždicový index drží nízké časy i pro větší dotazovací okna.",
            "U největšího okna roste počet kandidátů, takže roste i čas indexovaného dotazu.",
        ],
        width=5.5,
        size=16,
    )
    add_chart(s, charts["cr"])

    s = slide("Co ukazuje škálování")
    add_bullets(
        s,
        [
            "Bez indexu je rozhodující hlavně počet bodů: každý dotaz musí projít vše.",
            "U indexu je rozhodující počet kandidátů v dotčených dlaždicích nebo bboxu.",
            "Čím menší část prostoru dotaz zasáhne, tím větší je přínos indexu.",
            "Index má smysl hlavně při opakovaných dotazech nad stejným datasetem.",
        ],
        width=7.7,
        size=18,
    )
    add_note(s, "Obecné pravidlo:\nindex se vyplatí, když sníží počet přesně testovaných bodů.", left=8.35, top=2.05, width=3.8, height=2.8)

    s = slide("Radius query")
    add_bullets(
        s,
        [
            "Dotaz: najít všechny body do zadané vzdálenosti od středu.",
            "Nejdřív se použije bbox kružnice jako hrubý filtr.",
            "Potom musí následovat přesný test eukleidovské vzdálenosti.",
            "Poloměry v benchmarku: 5 km, 10 km a 25 km.",
        ],
        width=5.5,
        size=16,
    )
    add_chart(s, charts["radius"])

    s = slide("Polygon query")
    add_bullets(
        s,
        [
            "Dotaz: spočítat body ležící uvnitř nepravidelného polygonu kraje.",
            "Index najde kandidáty podle bounding boxu polygonu.",
            "Finální výsledek určuje až přesný test vůči skutečné hranici kraje.",
            "Použito 14 skutečných krajů České republiky.",
        ],
        width=5.5,
        size=16,
    )
    add_chart(s, charts["counts"])

    s = slide("Benchmark: polygon query", "14 krajů ČR, souhrnný čas podle velikosti datasetu")
    add_bullets(
        s,
        [
            "U polygon query je přesné ověření kandidátů dražší než u bbox dotazu.",
            "Pro 50 000 bodů je nejrychlejší R-tree/STRtree, dlaždicový index je druhý.",
            "U menších datasetů režie indexu může převážit nad přínosem.",
        ],
        width=5.5,
        size=16,
    )
    add_chart(s, charts["polygon"])

    s = slide("Radius vs. polygon query")
    add_bullets(
        s,
        [
            "Radius query: bbox kružnice je poměrně těsný filtr, přesný test je jednoduchá vzdálenost.",
            "Polygon query: bbox kraje může obsahovat mnoho bodů mimo skutečný tvar.",
            "Proto je u polygonů důležité, kolik falešných kandidátů projde hrubým filtrem.",
            "Různé typy dotazů mohou zvýhodnit jiný index.",
        ],
        width=7.8,
        size=18,
    )
    add_note(s, "Index urychluje kandidáty.\nSprávnost zajišťuje přesný geometrický test.", left=8.45, top=2.0, width=3.8, height=2.6)

    s = slide("Závěr")
    add_bullets(
        s,
        [
            "Pro bodové bbox a radius dotazy vyšel v projektu nejlépe jednoduchý dlaždicový index.",
            "Lineární průchod je dobrý jako referenční metoda, ale špatně škáluje.",
            "R-tree/STRtree je obecnější a u polygon query nad 50 000 body byl nejrychlejší.",
            "EPSG:5514 zjednodušuje interpretaci a umožňuje korektní metrické výpočty.",
            "Nejdůležitější praktický závěr: index je hrubý filtr, přesný geometrický test je stále nutný.",
        ],
        width=10.8,
        size=18,
    )

    prs.save(OUTPUT)


def main() -> None:
    PRESENTATION.mkdir(parents=True, exist_ok=True)
    charts = make_charts()
    build_presentation(charts)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
